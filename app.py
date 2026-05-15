import os
import re
import math
import streamlit as st
import pdfplumber
import docx
import google.generativeai as genai
from dotenv import load_dotenv
from fpdf import FPDF

load_dotenv()
genai.configure(api_key=os.getenv("GEMINI_API_KEY"))

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="HireIQ — AI Resume Screener",
    page_icon="🧠",
    layout="wide",
)

# ── CSS ───────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap');

*, *::before, *::after { box-sizing: border-box; }

[data-testid="stAppViewContainer"] {
    background: #06091a;
    font-family: 'Inter', sans-serif;
}
[data-testid="stHeader"] { background: transparent; }
[data-testid="stMain"] > div { padding-top: 0 !important; }
#MainMenu, footer { visibility: hidden; }

/* ── Hero ── */
.hero {
    position: relative;
    overflow: hidden;
    background: linear-gradient(160deg, #0b0f2a 0%, #130d35 55%, #0b0f2a 100%);
    border-bottom: 1px solid rgba(99,102,241,0.15);
    padding: 3rem 2.5rem 2.5rem;
    margin: -4rem -4rem 2.5rem;
}
.hero-glow {
    position: absolute;
    top: -80px; right: -80px;
    width: 420px; height: 420px;
    background: radial-gradient(circle, rgba(99,102,241,0.18) 0%, transparent 65%);
    border-radius: 50%;
    pointer-events: none;
}
.hero-glow-2 {
    position: absolute;
    bottom: -100px; left: 30%;
    width: 300px; height: 300px;
    background: radial-gradient(circle, rgba(139,92,246,0.1) 0%, transparent 65%);
    border-radius: 50%;
    pointer-events: none;
}
.hero-title {
    font-size: 3rem;
    font-weight: 900;
    letter-spacing: -0.03em;
    background: linear-gradient(135deg, #c7d2fe 0%, #a5b4fc 35%, #818cf8 70%, #6366f1 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
    margin: 0 0 0.35rem;
    line-height: 1.1;
}
.hero-subtitle {
    color: #475569;
    font-size: 0.95rem;
    font-weight: 500;
    letter-spacing: 0.03em;
    margin: 0 0 1.2rem;
}
.hero-badges {
    display: flex;
    flex-wrap: wrap;
    gap: 8px;
}
.hbadge {
    background: rgba(99,102,241,0.1);
    border: 1px solid rgba(99,102,241,0.25);
    color: #a5b4fc;
    font-size: 0.68rem;
    font-weight: 700;
    padding: 4px 12px;
    border-radius: 100px;
    letter-spacing: 0.07em;
    text-transform: uppercase;
}

/* ── Input labels ── */
.field-label {
    font-size: 0.72rem;
    font-weight: 700;
    color: #475569;
    letter-spacing: 0.1em;
    text-transform: uppercase;
    margin-bottom: 6px;
}

/* ── Score gauge wrapper ── */
.gauge-card {
    background: linear-gradient(135deg, #0d1130 0%, #160f38 100%);
    border: 1px solid rgba(99,102,241,0.2);
    border-radius: 18px;
    padding: 1.6rem 1.8rem;
    display: flex;
    align-items: center;
    gap: 1.6rem;
    margin-bottom: 1.4rem;
    position: relative;
    overflow: hidden;
}
.gauge-card::after {
    content: '';
    position: absolute;
    top: 0; left: 0; right: 0;
    height: 1px;
    background: linear-gradient(90deg, transparent, rgba(99,102,241,0.4), transparent);
}
.gauge-meta { flex: 1; }
.gauge-meta-label {
    font-size: 0.72rem;
    font-weight: 700;
    color: #475569;
    text-transform: uppercase;
    letter-spacing: 0.1em;
    margin-bottom: 6px;
}
.gauge-verdict {
    font-size: 1.5rem;
    font-weight: 800;
    margin-bottom: 6px;
    letter-spacing: -0.02em;
}
.gauge-desc {
    font-size: 0.82rem;
    color: #64748b;
    line-height: 1.5;
    max-width: 260px;
}

/* ── Two-up grid ── */
.two-up {
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 1rem;
    margin-bottom: 1rem;
}
@media (max-width: 700px) { .two-up { grid-template-columns: 1fr; } }

/* ── Section card ── */
.scard {
    background: #0c1128;
    border: 1px solid #1a2240;
    border-radius: 14px;
    padding: 1.2rem 1.4rem;
}
.scard-title {
    font-size: 0.7rem;
    font-weight: 800;
    letter-spacing: 0.1em;
    text-transform: uppercase;
    display: flex;
    align-items: center;
    gap: 8px;
    margin-bottom: 1rem;
    padding-bottom: 0.7rem;
    border-bottom: 1px solid #1a2240;
}
.scard-icon {
    width: 22px; height: 22px;
    border-radius: 6px;
    display: flex; align-items: center; justify-content: center;
    font-size: 0.75rem;
    flex-shrink: 0;
}
.scard-icon-green { background: rgba(16,185,129,0.15); }
.scard-icon-red   { background: rgba(239,68,68,0.15); }
.scard-title-green { color: #34d399; }
.scard-title-red   { color: #f87171; }

/* ── Keyword chips ── */
.chips { display: flex; flex-wrap: wrap; gap: 6px; }
.chip {
    padding: 4px 11px;
    border-radius: 100px;
    font-size: 0.75rem;
    font-weight: 600;
    line-height: 1.4;
    white-space: nowrap;
}
.chip-red {
    background: rgba(239,68,68,0.1);
    border: 1px solid rgba(239,68,68,0.25);
    color: #fca5a5;
}
.chip-green {
    background: rgba(16,185,129,0.1);
    border: 1px solid rgba(16,185,129,0.25);
    color: #6ee7b7;
}

/* ── Strength rows ── */
.strength-row {
    display: flex;
    gap: 10px;
    align-items: flex-start;
    padding: 10px 0;
    border-bottom: 1px solid #111827;
}
.strength-row:last-child { border-bottom: none; }
.strength-dot {
    width: 7px; height: 7px;
    border-radius: 50%;
    background: #34d399;
    margin-top: 6px;
    flex-shrink: 0;
}
.strength-title { font-size: 0.86rem; font-weight: 700; color: #e2e8f0; }
.strength-body  { font-size: 0.8rem; color: #64748b; margin-top: 2px; }

/* ── Resume file header ── */
.file-header {
    display: flex;
    align-items: center;
    gap: 12px;
    background: rgba(99,102,241,0.07);
    border: 1px solid rgba(99,102,241,0.15);
    border-radius: 12px;
    padding: 12px 16px;
    margin: 1.8rem 0 1.2rem;
}
.file-icon {
    width: 38px; height: 38px;
    background: rgba(99,102,241,0.15);
    border-radius: 9px;
    display: flex; align-items: center; justify-content: center;
    font-size: 1.1rem;
    flex-shrink: 0;
}
.file-name { font-size: 1rem; font-weight: 700; color: #e2e8f0; }
.file-sub  { font-size: 0.75rem; color: #475569; margin-top: 2px; }

/* ── Agent accordion override ── */
[data-testid="stExpander"] {
    background: #0c1128 !important;
    border: 1px solid #1a2240 !important;
    border-radius: 12px !important;
}
[data-testid="stExpander"] summary {
    font-size: 0.82rem !important;
    font-weight: 600 !important;
    color: #64748b !important;
    letter-spacing: 0.03em !important;
}
[data-testid="stExpander"] summary:hover { color: #94a3b8 !important; }

/* ── Explanation card ── */
.explain-card {
    background: linear-gradient(135deg, rgba(99,102,241,0.07), rgba(139,92,246,0.07));
    border: 1px solid rgba(99,102,241,0.18);
    border-radius: 14px;
    padding: 1.3rem 1.5rem;
    margin-bottom: 1rem;
}
.explain-label {
    font-size: 0.7rem;
    font-weight: 800;
    color: #6366f1;
    letter-spacing: 0.1em;
    text-transform: uppercase;
    margin-bottom: 0.7rem;
}
.explain-text {
    font-size: 0.9rem;
    color: #c7d2fe;
    line-height: 1.8;
}

/* ── Buttons ── */
[data-testid="stButton"] > button[kind="primary"] {
    background: linear-gradient(135deg, #4338ca, #6366f1) !important;
    border: none !important;
    border-radius: 12px !important;
    font-family: 'Inter', sans-serif !important;
    font-weight: 700 !important;
    font-size: 0.95rem !important;
    letter-spacing: 0.02em !important;
    padding: 0.75rem 2rem !important;
    box-shadow: 0 4px 20px rgba(99,102,241,0.35) !important;
    transition: all 0.2s ease !important;
}
[data-testid="stButton"] > button[kind="primary"]:hover {
    box-shadow: 0 6px 28px rgba(99,102,241,0.55) !important;
    transform: translateY(-1px) !important;
}
[data-testid="stDownloadButton"] > button {
    background: #0c1128 !important;
    border: 1px solid rgba(99,102,241,0.3) !important;
    color: #a5b4fc !important;
    border-radius: 10px !important;
    font-weight: 600 !important;
    transition: all 0.2s !important;
}
[data-testid="stDownloadButton"] > button:hover {
    background: rgba(99,102,241,0.1) !important;
    border-color: rgba(99,102,241,0.5) !important;
}

/* ── Inputs ── */
div[data-testid="stTextArea"] textarea {
    background: #0c1128 !important;
    color: #cbd5e1 !important;
    border: 1px solid #1a2240 !important;
    border-radius: 10px !important;
    font-size: 0.84rem !important;
    font-family: 'Inter', sans-serif !important;
    line-height: 1.65 !important;
}
div[data-testid="stTextArea"] textarea:focus {
    border-color: rgba(99,102,241,0.5) !important;
    box-shadow: 0 0 0 3px rgba(99,102,241,0.12) !important;
}
[data-testid="stFileUploaderDropzone"] {
    background: #0c1128 !important;
    border: 2px dashed #1a2240 !important;
    border-radius: 12px !important;
    transition: border-color 0.2s !important;
}
[data-testid="stFileUploaderDropzone"]:hover {
    border-color: rgba(99,102,241,0.4) !important;
}

/* ── Divider ── */
hr[data-testid="stDivider"] {
    border-color: #111827 !important;
    margin: 1.8rem 0 !important;
}

/* ── Warning / info ── */
[data-testid="stAlert"] {
    background: #0c1128 !important;
    border: 1px solid #1a2240 !important;
    border-radius: 10px !important;
}

/* ── Footer ── */
.footer {
    text-align: center;
    padding: 3rem 0 1.5rem;
    color: #1e293b;
    font-size: 0.78rem;
    font-weight: 500;
}
.footer a { color: #4338ca; text-decoration: none; }
.footer a:hover { color: #818cf8; }
</style>
""", unsafe_allow_html=True)

# ── Resume text extraction ────────────────────────────────────────────────────
def extract_text_from_pdf(file):
    with pdfplumber.open(file) as pdf:
        return "\n".join([p.extract_text() or "" for p in pdf.pages])

def extract_text_from_docx(file):
    doc = docx.Document(file)
    return "\n".join([p.text for p in doc.paragraphs])

def get_resume_text(uploaded_file):
    if uploaded_file.name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif uploaded_file.name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    return "Unsupported file format."

# ── Gemini helper ─────────────────────────────────────────────────────────────
def gemini_prompt(system_prompt, user_input):
    model = genai.GenerativeModel("gemini-1.5-flash")
    response = model.generate_content(
        [{"role": "user", "parts": [system_prompt + "\n\n" + user_input]}]
    )
    return response.text

# ── Score parser ──────────────────────────────────────────────────────────────
def extract_score(text):
    for pat in [
        r"SCORE\s*:\s*(\d{1,3})\s*/\s*100",
        r"(\d{1,3})\s*/\s*100",
        r"[Ss]core[:\s]+(\d{1,3})",
        r"(\d{1,3})\s*out\s*of\s*100",
        r"(\d{1,3})\s*%",
    ]:
        m = re.search(pat, text)
        if m:
            val = int(m.group(1))
            if 0 <= val <= 100:
                return val
    return None

# ── 5-Agent pipeline ──────────────────────────────────────────────────────────
def agentic_resume_screening(resume_text, job_description):
    info = gemini_prompt(
        "You are a resume parser. Extract all structured candidate details: "
        "skills, years of experience, education, certifications, and key achievements. "
        "Be thorough and organised.",
        resume_text,
    )
    match = gemini_prompt(
        "You are a senior recruiter. Match the resume to the job description. "
        "Start your response with a line in EXACTLY this format: 'SCORE: XX/100' "
        "(replace XX with the numeric score). "
        "Then list the key matching skills and experiences.",
        f"Resume Info:\n{info}\n\nJob Description:\n{job_description}",
    )
    explain = gemini_prompt(
        "You are an HR business partner. Write a 3-5 sentence plain-language summary "
        "of how well this candidate fits the role, suitable for a hiring manager with "
        "no technical background.",
        match,
    )
    missing = gemini_prompt(
        "You are a skills-gap analyst. List ONLY the skills, technologies, and keywords "
        "that appear in the job description but are ABSENT from the resume. "
        "Return a clean bullet list (• item). Be concise — no explanations needed.",
        f"Resume Info:\n{info}\n\nJob Description:\n{job_description}",
    )
    strengths = gemini_prompt(
        "You are a career coach. Identify the top 3-5 strengths of this resume "
        "relative to this specific job description. "
        "Return a bullet list where each point is: • Strength — one-line reason why it matters for this role.",
        f"Resume Info:\n{info}\n\nJob Description:\n{job_description}",
    )
    return info, match, explain, missing, strengths

# ── HTML renderers ────────────────────────────────────────────────────────────
def score_gauge_html(score, color):
    r = 50
    circ = 2 * math.pi * r
    offset = circ * (1 - score / 100)
    verdict, desc = (
        ("Strong Match", "This candidate aligns well with the role requirements.")
        if score >= 70 else
        ("Partial Match", "Candidate meets some requirements but has notable gaps.")
        if score >= 50 else
        ("Weak Match", "Significant skill gaps relative to the job description.")
    )
    return f"""
    <div class="gauge-card">
        <svg width="120" height="120" viewBox="0 0 120 120" style="flex-shrink:0;">
            <circle cx="60" cy="60" r="{r}" fill="none" stroke="#111827" stroke-width="8"/>
            <circle cx="60" cy="60" r="{r}" fill="none" stroke="{color}" stroke-width="8"
                    stroke-dasharray="{circ:.2f}" stroke-dashoffset="{offset:.2f}"
                    stroke-linecap="round" transform="rotate(-90 60 60)"/>
            <text x="60" y="54" text-anchor="middle" fill="{color}"
                  font-size="26" font-weight="800" font-family="Inter,sans-serif">{score}</text>
            <text x="60" y="70" text-anchor="middle" fill="#334155"
                  font-size="10" font-family="Inter,sans-serif">/ 100</text>
        </svg>
        <div class="gauge-meta">
            <div class="gauge-meta-label">Match Score</div>
            <div class="gauge-verdict" style="color:{color};">{verdict}</div>
            <div class="gauge-desc">{desc}</div>
        </div>
    </div>"""

def _parse_bullets(text):
    return [
        l.strip().lstrip("•·▪*-").strip()
        for l in text.splitlines()
        if l.strip() and len(l.strip().lstrip("•·▪*-").strip()) > 1
    ]

def keyword_chips_html(text):
    items = _parse_bullets(text)
    if not items:
        return f'<p style="color:#475569;font-size:0.83rem;">{text}</p>'
    chips = "".join(f'<span class="chip chip-red">{i}</span>' for i in items if len(i) < 50)
    return f'<div class="chips">{chips}</div>'

def strength_cards_html(text):
    items = _parse_bullets(text)
    if not items:
        return f'<p style="color:#475569;font-size:0.83rem;">{text}</p>'
    rows = []
    for item in items:
        for sep in (" — ", " - ", ": "):
            if sep in item:
                title, body = item.split(sep, 1)
                rows.append(f"""
                <div class="strength-row">
                    <div class="strength-dot"></div>
                    <div>
                        <div class="strength-title">{title.strip()}</div>
                        <div class="strength-body">{body.strip()}</div>
                    </div>
                </div>""")
                break
        else:
            rows.append(f"""
            <div class="strength-row">
                <div class="strength-dot"></div>
                <div><div class="strength-title">{item}</div></div>
            </div>""")
    return "".join(rows)

# ── PDF export ────────────────────────────────────────────────────────────────
def generate_pdf_report(filename, extracted_info, match_report, explanation,
                        missing_keywords, top_strengths, score):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(99, 102, 241)
    pdf.cell(0, 14, "HireIQ  AI Resume Screener", ln=True, align="C")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 7, f"Resume: {filename}", ln=True, align="C")
    pdf.ln(4)
    if score is not None:
        clr = (34, 197, 94) if score >= 70 else (245, 158, 11) if score >= 50 else (239, 68, 68)
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(*clr)
        pdf.cell(0, 10, f"Match Score: {score}/100", ln=True)
        pdf.ln(2)
    for title, body in [
        ("Agent 1 — Extracted Resume Info", extracted_info),
        ("Agent 2 — Match Report", match_report),
        ("Agent 3 — HR-Friendly Explanation", explanation),
        ("Agent 4 — Missing Keywords", missing_keywords),
        ("Agent 5 — Top Strengths", top_strengths),
    ]:
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(30, 41, 59)
        pdf.set_fill_color(241, 245, 249)
        pdf.cell(0, 9, f"  {title}", ln=True, fill=True)
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(51, 65, 85)
        pdf.multi_cell(0, 5.5, body.encode("latin-1", "replace").decode("latin-1"))
        pdf.ln(3)
    return bytes(pdf.output())

# ── Sample JD ─────────────────────────────────────────────────────────────────
SAMPLE_JD = """\
Software Engineer — Full Stack
Company: TechCorp Inc.

We are looking for a passionate Software Engineer with hands-on experience in:
• Python, JavaScript, TypeScript
• React, Node.js, REST APIs and GraphQL
• SQL (PostgreSQL / MySQL) and NoSQL databases (MongoDB, Redis)
• Cloud platforms: AWS or GCP
• Docker, Kubernetes, CI/CD pipelines (GitHub Actions / Jenkins)
• Git version control, Agile / Scrum methodology

Responsibilities:
- Design and develop scalable, production-ready web applications
- Collaborate with cross-functional teams (Product, Design, Data)
- Write clean, well-tested code; participate in code reviews
- Improve system reliability through monitoring and observability

Requirements:
- B.S. in Computer Science or related field
- 2+ years of professional software development experience
- Strong problem-solving and communication skills
- Experience with machine learning or data pipelines is a plus
"""

# ── Hero ──────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="hero">
    <div class="hero-glow"></div>
    <div class="hero-glow-2"></div>
    <div class="hero-title">🧠 HireIQ</div>
    <div class="hero-subtitle">AI-Powered Resume Screener &nbsp;·&nbsp; 5-Agent Gemini Flash Pipeline</div>
    <div class="hero-badges">
        <span class="hbadge">Gemini Flash</span>
        <span class="hbadge">5 AI Agents</span>
        <span class="hbadge">PDF Export</span>
        <span class="hbadge">Multi-Resume</span>
    </div>
</div>
""", unsafe_allow_html=True)

# ── Inputs ────────────────────────────────────────────────────────────────────
col_jd, col_up = st.columns([1, 1], gap="large")

with col_jd:
    st.markdown('<div class="field-label">📋 Job Description</div>', unsafe_allow_html=True)
    job_description = st.text_area(
        label="jd_hidden",
        value=SAMPLE_JD,
        height=300,
        label_visibility="collapsed",
        help="A sample JD is pre-filled — upload a resume and click Analyze to test right away.",
    )

with col_up:
    st.markdown('<div class="field-label">📂 Upload Resume(s) — PDF or DOCX</div>', unsafe_allow_html=True)
    uploaded_resumes = st.file_uploader(
        label="ul_hidden",
        type=["pdf", "docx"],
        accept_multiple_files=True,
        label_visibility="collapsed",
    )
    st.caption("Upload multiple resumes to screen several candidates at once.")

st.divider()

# ── Analyze ───────────────────────────────────────────────────────────────────
if st.button("🔍  Analyze Resumes", type="primary", use_container_width=True):
    if not job_description.strip() or not uploaded_resumes:
        st.warning("Please upload at least one resume and provide a job description.")
    else:
        for file in uploaded_resumes:
            # File header
            st.markdown(f"""
            <div class="file-header">
                <div class="file-icon">📄</div>
                <div>
                    <div class="file-name">{file.name}</div>
                    <div class="file-sub">Running 5-agent analysis pipeline</div>
                </div>
            </div>""", unsafe_allow_html=True)

            resume_text = get_resume_text(file)

            with st.spinner("Running agents…"):
                extracted_info, match_report, explanation, missing_keywords, top_strengths = (
                    agentic_resume_screening(resume_text, job_description)
                )

            score = extract_score(match_report)
            score_color = (
                "#10b981" if score is not None and score >= 70 else
                "#f59e0b" if score is not None and score >= 50 else
                "#ef4444"
            )

            # Score gauge
            if score is not None:
                st.markdown(score_gauge_html(score, score_color), unsafe_allow_html=True)
            else:
                st.info("Could not parse a numeric score — check the Match Report below.")

            # Strengths + Missing Keywords
            st.markdown("""<div class="two-up">""", unsafe_allow_html=True)

            col_s, col_m = st.columns(2, gap="medium")

            with col_s:
                st.markdown(f"""
                <div class="scard">
                    <div class="scard-title scard-title-green">
                        <div class="scard-icon scard-icon-green">✦</div>
                        Top Strengths
                    </div>
                    {strength_cards_html(top_strengths)}
                </div>""", unsafe_allow_html=True)

            with col_m:
                st.markdown(f"""
                <div class="scard">
                    <div class="scard-title scard-title-red">
                        <div class="scard-icon scard-icon-red">⚠</div>
                        Missing Keywords
                    </div>
                    {keyword_chips_html(missing_keywords)}
                </div>""", unsafe_allow_html=True)

            st.markdown("</div>", unsafe_allow_html=True)

            # HR Explanation
            safe_explain = explanation.replace("<", "&lt;").replace(">", "&gt;")
            st.markdown(f"""
            <div class="explain-card">
                <div class="explain-label">💬 HR-Friendly Explanation</div>
                <div class="explain-text">{safe_explain}</div>
            </div>""", unsafe_allow_html=True)

            # Raw agent outputs
            with st.expander("🕵️ Agent 1 — Extracted Resume Info"):
                st.text(extracted_info)
            with st.expander("🎯 Agent 2 — Match Report"):
                st.text(match_report)

            # PDF download
            pdf_bytes = generate_pdf_report(
                file.name, extracted_info, match_report, explanation,
                missing_keywords, top_strengths, score,
            )
            stem = file.name.rsplit(".", 1)[0]
            st.download_button(
                label="📥  Download Full Report (PDF)",
                data=pdf_bytes,
                file_name=f"HireIQ_Report_{stem}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )

            st.divider()

# ── Footer ────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="footer">
    Built by <a href="https://github.com/vinay23is" target="_blank">Vinay Dodla</a>
    &nbsp;·&nbsp;
    Powered by Gemini Flash &amp; Streamlit
    &nbsp;·&nbsp;
    <a href="https://github.com/vinay23is/smart-resume-screener" target="_blank">GitHub</a>
</div>
""", unsafe_allow_html=True)
